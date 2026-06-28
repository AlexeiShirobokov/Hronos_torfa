import unittest
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
import build_xlsx, analytics
from tests.test_analytics import _make_csv, SCRATCH


class TestBuildXlsx(unittest.TestCase):
    def test_sheets_appended(self):
        csv = SCRATCH / "consolidated.csv"; _make_csv(csv)
        book = SCRATCH / "reestr.xlsx"
        pd.DataFrame({"Подразделение": ["Эрел"]}).to_excel(
            book, index=False, sheet_name="Сводный_Реестр")
        aggr = analytics.compute(csv, "2026-06-28")
        build_xlsx.write_sheets(aggr, book)
        wb = load_workbook(book)
        for s in ("Свод_подразделения", "Свод_даты", "ABC_водители", "Парк_марки",
                  "Парк_инв", "Машины_дни", "Машины_смены", "Машины_час_Торф"):
            self.assertIn(s, wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
