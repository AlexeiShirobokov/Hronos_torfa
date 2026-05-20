"""Собрать консолидированный Excel из вложений в input/mail_attachments.

Берёт самые поздние версии каждого «базового» имени, объединяет их «Реестр»
в один лист 'Сводный_Реестр' и сохраняет в output/<имя>_<дата>.xlsx.

Используется как run_daily.py-шаг 2.
"""
from __future__ import annotations
import json, os, re, sys
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
ATT = BASE / "input" / "mail_attachments"
OUT = BASE / "output"
LOGS = BASE / "logs"

ALL_COLS = [
    "Подразделение", "Дата. Факт", "Дата выдачи наряд-задания", "Смена",
    "Ф.И.О. Ответственного", "Ф.И.О. Машиниста экскватора",
    "Ф.И.О. водителя самосвала", "Время", "Блок",
    "Марка погрузочной единицы", "Инв. № погрузочной единицы",
    "Марка транспортировочной единицы", "Инв. № транспортировочной единицы",
    "Количство машин, шт", "Обьем работ, м3", "Обьем кузова,м3",
    "Ф.И.О. моториста п/п 1", "Ф.И.О. моториста п/п 2",
    "Марка промывочного прибора", "Инв. № промывочного прибора",
    "Передел", "Откатка, м", "Примечание",
]


def _clean(v):
    if v is None: return None
    if isinstance(v, float) and v != v: return None
    try:
        if pd.isna(v): return None
    except (TypeError, ValueError): pass
    return v


def main() -> int:
    if not ATT.exists():
        print(f"[ERR] нет папки с вложениями: {ATT}"); return 2
    files = sorted(p.name for p in ATT.iterdir() if p.is_file())
    if not files:
        print(f"[ERR] папка вложений пуста: {ATT}"); return 3

    groups: dict[str, list[str]] = {}
    for f in files:
        base = re.sub(r"__\d+(?=\.[^.]+$)", "", f)
        groups.setdefault(base, []).append(f)

    def ver(f):
        m = re.search(r"__(\d+)(?=\.[^.]+$)", f)
        return int(m.group(1)) if m else 0

    latest = {b: max(v, key=ver) for b, v in groups.items()}

    frames = []
    used = []
    for base, fname in latest.items():
        try:
            df = pd.read_excel(ATT / fname, sheet_name="Реестр",
                               engine="openpyxl").dropna(how="all")
            frames.append(df)
            used.append({"base": base, "file": fname, "rows": len(df)})
        except Exception as e:
            used.append({"base": base, "file": fname, "rows": 0,
                         "error": f"{type(e).__name__}: {e}"})
    if not frames:
        print("[ERR] не удалось прочитать ни одного 'Реестр'"); return 4

    full = pd.concat(frames, ignore_index=True)
    for c in ALL_COLS:
        if c not in full.columns: full[c] = pd.NA
    extras = [c for c in full.columns if c not in ALL_COLS]
    full = full[ALL_COLS + extras]
    full["Подразделение"] = full["Подразделение"].apply(
        lambda x: x.strip() if isinstance(x, str) else x
    )

    meta_file = LOGS / "last_fetch.meta"
    report_date = datetime.now().strftime("%Y-%m-%d")
    if meta_file.exists():
        for line in meta_file.read_text(encoding="utf-8").splitlines():
            if line.startswith("date="):
                report_date = line.split("=", 1)[1].strip()

    OUT.mkdir(parents=True, exist_ok=True)
    out_path = OUT / f"Хронометраж_транспортировки_торфов_{report_date}.xlsx"

    wb = Workbook(); ws = wb.active; ws.title = "Сводный_Реестр"
    ws.append(list(full.columns))
    for row in full.itertuples(index=False, name=None):
        ws.append([_clean(v) for v in row])
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="305496")
    for c, _ in enumerate(full.columns, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font; cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    for c, name in enumerate(full.columns, 1):
        L = max(12, min(40, int(len(str(name)) * 1.1) + 2))
        ws.column_dimensions[get_column_letter(c)].width = L
    ws.freeze_panes = "A2"
    if len(full):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(full.columns))}{len(full)+1}"
    wb.save(out_path)

    # промежуточный csv для PDF-шага
    mid = LOGS / "consolidated.csv"
    full.to_csv(mid, index=False)

    info = {
        "report_date": report_date,
        "rows": int(len(full)),
        "files_used": used,
        "out": str(out_path),
        "csv": str(mid),
        "ts": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    (LOGS / "last_consolidate.json").write_text(
        json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"ok": True, **info}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
