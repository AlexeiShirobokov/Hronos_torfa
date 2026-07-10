from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import (
    ALL_COLS,
    CHECK_SHEET,
    CRITICAL_COLS,
    DEFAULT_REQUIRED_UNITS,
    EXCLUDED_OUTPUT_COLS,
    FILES_SHEET,
    FRESHNESS_SHEET,
    OUTPUT_COLS,
    PIVOT_EXPECTED_SHEET,
    REGISTRY_SHEET,
    DEFAULT_PIVOT_PEREDELS,
    SAND_PEREDL,
    SAND_TO_WAREHOUSE_PEREDL,
    canonicalize_peredel,
    canonicalize_columns,
    norm_text,
)

LOG = logging.getLogger("hronos.consolidate")
EXCEL_DATE_FORMAT = "[$-419]d\\ mmm;@"
EXCEL_EXTS = {".xlsx", ".xlsb"}
VERSION_RE = re.compile(r"__(\d+)(?=\.[^.]+$)")


@dataclass
class FileReport:
    path: str
    base_name: str
    rows: int = 0
    volume: float = 0.0
    volumes_by_unit: dict[str, float] = field(default_factory=dict)
    units: list[str] = field(default_factory=list)
    max_fact_date: str | None = None
    max_issue_date: str | None = None
    error: str | None = None


@dataclass
class ConsolidationResult:
    workbook: Path
    metadata: Path
    rows: int
    files_used: list[FileReport]
    volume_check: pd.DataFrame
    freshness_check: pd.DataFrame
    pivot_expected: pd.DataFrame


def discover_latest_files(input_dir: Path) -> list[Path]:
    files = [
        p for p in sorted(input_dir.iterdir())
        if p.is_file() and p.suffix.lower() in EXCEL_EXTS and not p.name.startswith("~$")
        and "хронометраж" in p.name.casefold()
    ]
    groups: dict[str, list[Path]] = {}
    for path in files:
        groups.setdefault(_base_name(path.name), []).append(path)
    return [max(paths, key=lambda p: (_version(p), p.stat().st_mtime)) for paths in groups.values()]


def consolidate(
    input_dir: Path,
    output_dir: Path,
    report_date: date | None = None,
    sheet_name: str = "Реестр",
    pivot_peredels: Iterable[str] | None = None,
    freshness_date: date | None = None,
    required_units: Iterable[str] | None = None,
) -> ConsolidationResult:
    input_dir = input_dir.resolve()
    output_dir = output_dir.resolve()
    report_date = report_date or date.today()

    if not input_dir.exists():
        raise FileNotFoundError(f"нет папки входных файлов: {input_dir}")

    files = discover_latest_files(input_dir)
    if not files:
        raise FileNotFoundError(f"не найдены .xlsx/.xlsb в {input_dir}")

    LOG.info("найдено файлов к обработке: %s", len(files))
    frames: list[pd.DataFrame] = []
    reports: list[FileReport] = []

    for path in files:
        report = FileReport(path=str(path), base_name=_base_name(path.name))
        try:
            df = read_registry(path, sheet_name=sheet_name)
            df = normalize_registry(df, path)
            report.rows = int(len(df))
            report.volume = _volume_sum(df)
            report.volumes_by_unit = _volumes_by_unit(df)
            report.units = sorted(report.volumes_by_unit)
            report.max_fact_date = _max_date_iso(df["Дата. Факт"])
            report.max_issue_date = _max_date_iso(df["Дата выдачи наряд-задания"])
            frames.append(df)
        except Exception as exc:
            report.error = f"{type(exc).__name__}: {exc}"
            LOG.error("не удалось прочитать %s: %s", path.name, report.error)
        reports.append(report)

    if not frames:
        raise RuntimeError("не удалось прочитать ни одного реестра")

    registry = pd.concat(frames, ignore_index=True)
    registry = _order_columns(registry)
    volume_check = build_volume_check(reports, registry)
    freshness_check = build_freshness_check(reports, freshness_date, required_units)
    pivot_expected = build_pivot_expected(registry, report_date, pivot_peredels)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = output_dir / f"Хронометраж_по_подразделениям_{report_date:%Y-%m-%d}_{ts}.xlsx"
    metadata = workbook.with_suffix(".json")

    write_workbook(workbook, registry, volume_check, freshness_check, reports, pivot_expected)
    metadata.write_text(
        json.dumps(
            {
                "ok": True,
                "report_date": report_date.isoformat(),
                "rows": int(len(registry)),
                "workbook": str(workbook),
                "files_used": [r.__dict__ for r in reports],
                "volume_check": volume_check.to_dict(orient="records"),
                "freshness_date": _freshness_target(reports, freshness_date).isoformat(),
                "freshness_check": freshness_check.to_dict(orient="records"),
                "pivot_peredels": list(pivot_peredels or DEFAULT_PIVOT_PEREDELS),
                "pivot_expected": pivot_expected.to_dict(orient="records"),
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    LOG.info("консолидация готова: %s строк, файл %s", len(registry), workbook.name)
    return ConsolidationResult(workbook, metadata, len(registry), reports, volume_check, freshness_check, pivot_expected)


def read_registry(path: Path, sheet_name: str = "Реестр") -> pd.DataFrame:
    engine = "pyxlsb" if path.suffix.lower() == ".xlsb" else "openpyxl"
    xls = pd.ExcelFile(path, engine=engine)
    selected = _select_sheet(xls.sheet_names, sheet_name)
    if selected is None:
        raise ValueError(f"нет листа «{sheet_name}»")
    return pd.read_excel(xls, sheet_name=selected)


def normalize_registry(df: pd.DataFrame, path: Path) -> pd.DataFrame:
    df = df.dropna(how="all").copy()
    df.columns = canonicalize_columns(df.columns)
    df = _merge_duplicate_alias_columns(df)

    missing_critical = [col for col in CRITICAL_COLS if col not in df.columns]
    if missing_critical:
        raise ValueError("нет обязательных колонок реестра: " + ", ".join(missing_critical))

    for col in ALL_COLS:
        if col not in df.columns:
            df[col] = pd.NA

    if df["Подразделение"].isna().all():
        guessed = _unit_from_filename(path.name)
        if guessed:
            df["Подразделение"] = guessed

    df["Подразделение"] = df["Подразделение"].map(_clean_text_or_na)
    bad_units = {"", "nan", "<na>", "none", "итог", "итого"}
    mask_good = df["Подразделение"].notna() & ~df["Подразделение"].astype(str).str.casefold().isin(bad_units)
    df = df[mask_good].reset_index(drop=True)

    df["Передел"] = df["Передел"].map(lambda x: canonicalize_peredel(x) if pd.notna(x) else pd.NA)
    note_is_warehouse = df["Примечание"].map(norm_text).str.casefold().str.contains("склад", na=False)
    df.loc[(df["Передел"] == SAND_PEREDL) & note_is_warehouse, "Передел"] = SAND_TO_WAREHOUSE_PEREDL
    df["Время"] = df["Время"].map(norm_time)
    df["Дата. Факт"] = normalize_date_series(df["Дата. Факт"])
    df["Дата выдачи наряд-задания"] = normalize_date_series(df["Дата выдачи наряд-задания"])
    df["Обьем работ, м3"] = normalize_number_series(df["Обьем работ, м3"])
    df = df.drop(columns=[col for col in EXCLUDED_OUTPUT_COLS if col in df.columns], errors="ignore")
    return df


def build_volume_check(reports: Iterable[FileReport], registry: pd.DataFrame) -> pd.DataFrame:
    source: dict[str, float] = {}
    rows_by_unit: dict[str, int] = {}
    files_by_unit: dict[str, set[str]] = {}

    for report in reports:
        if report.error:
            continue
        for unit, volume in report.volumes_by_unit.items():
            source[unit] = source.get(unit, 0.0) + float(volume)
            files_by_unit.setdefault(unit, set()).add(Path(report.path).name)

    for unit, group in registry.groupby("Подразделение", dropna=False):
        rows_by_unit[str(unit)] = int(len(group))

    registry_totals = _volumes_by_unit(registry)
    units = sorted(set(source) | set(registry_totals))
    rows = []
    for unit in units:
        src = round(source.get(unit, 0.0), 3)
        dst = round(registry_totals.get(unit, 0.0), 3)
        delta = round(dst - src, 3)
        rows.append(
            {
                "Подразделение": unit,
                "Файлов": len(files_by_unit.get(unit, set())),
                "Строк": rows_by_unit.get(unit, 0),
                "Объем_источники": src,
                "Объем_реестр": dst,
                "Отклонение": delta,
                "Статус": "OK" if abs(delta) < 0.001 else "ПРОВЕРИТЬ",
            }
        )
    return pd.DataFrame(rows)


def build_freshness_check(
    reports: Iterable[FileReport],
    freshness_date: date | None = None,
    required_units: Iterable[str] | None = None,
) -> pd.DataFrame:
    reports = list(reports)
    target = _freshness_target(reports, freshness_date)
    units = [norm_text(unit) for unit in (required_units or DEFAULT_REQUIRED_UNITS)]
    rows = []
    for unit in units:
        unit_reports = [r for r in reports if unit in r.volumes_by_unit and not r.error]
        max_fact = max((r.max_fact_date for r in unit_reports if r.max_fact_date), default=None)
        max_issue = max((r.max_issue_date for r in unit_reports if r.max_issue_date), default=None)
        rows.append(
            {
                "Подразделение": unit,
                "Контрольная_дата": target.isoformat(),
                "Макс_Дата_Факт": max_fact,
                "Макс_Дата_выдачи_наряд-задания": max_issue,
                "Файлы": ", ".join(Path(r.path).name for r in unit_reports),
                "Статус": "OK" if max_fact and max_fact >= target.isoformat() else "НЕТ СВЕЖЕГО ФАЙЛА",
            }
        )
    return pd.DataFrame(rows)


def build_pivot_expected(
    registry: pd.DataFrame,
    report_date: date,
    pivot_peredels: Iterable[str] | None = None,
) -> pd.DataFrame:
    peredels = {canonicalize_peredel(x) for x in (pivot_peredels or DEFAULT_PIVOT_PEREDELS)}
    issue_dates = pd.to_datetime(registry["Дата выдачи наряд-задания"], errors="coerce").dt.date
    mask = (issue_dates == report_date) & registry["Передел"].map(canonicalize_peredel).isin(peredels)
    subset = registry[mask].copy()
    if subset.empty:
        return pd.DataFrame(columns=["Подразделение", "Ожидаемый_объем", "Переделы"])
    out = (
        subset.groupby("Подразделение", dropna=False)["Обьем работ, м3"]
        .sum()
        .reset_index()
        .rename(columns={"Обьем работ, м3": "Ожидаемый_объем"})
    )
    out["Ожидаемый_объем"] = out["Ожидаемый_объем"].round(3)
    out["Переделы"] = ", ".join(sorted(peredels))
    return out.sort_values("Подразделение").reset_index(drop=True)


def write_workbook(
    path: Path,
    registry: pd.DataFrame,
    volume_check: pd.DataFrame,
    freshness_check: pd.DataFrame,
    reports: list[FileReport],
    pivot_expected: pd.DataFrame,
) -> None:
    files_df = pd.DataFrame([r.__dict__ for r in reports])
    with pd.ExcelWriter(path, engine="openpyxl", date_format=EXCEL_DATE_FORMAT, datetime_format=EXCEL_DATE_FORMAT) as writer:
        registry.to_excel(writer, index=False, sheet_name=REGISTRY_SHEET)
        volume_check.to_excel(writer, index=False, sheet_name=CHECK_SHEET)
        freshness_check.to_excel(writer, index=False, sheet_name=FRESHNESS_SHEET)
        files_df.to_excel(writer, index=False, sheet_name=FILES_SHEET)
        pivot_expected.to_excel(writer, index=False, sheet_name=PIVOT_EXPECTED_SHEET)

    from openpyxl import load_workbook

    wb = load_workbook(path)
    for ws in wb.worksheets:
        _format_sheet(ws)
    for ws_name in (REGISTRY_SHEET,):
        ws = wb[ws_name]
        _format_date_columns(ws, ("Дата. Факт", "Дата выдачи наряд-задания"))
    wb.save(path)


def normalize_date_series(series: pd.Series) -> pd.Series:
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    numeric = pd.to_numeric(series, errors="coerce")
    numeric_mask = numeric.between(20_000, 80_000)
    if numeric_mask.any():
        parsed.loc[numeric_mask] = pd.to_datetime(
            numeric.loc[numeric_mask],
            unit="D",
            origin="1899-12-30",
            errors="coerce",
        )
    text_mask = ~numeric_mask
    parsed.loc[text_mask] = pd.to_datetime(series.loc[text_mask], errors="coerce")
    fallback_mask = parsed.isna() & series.notna()
    if fallback_mask.any():
        parsed.loc[fallback_mask] = pd.to_datetime(series.loc[fallback_mask], errors="coerce", dayfirst=True)
    return parsed.dt.normalize()


def normalize_number_series(series: pd.Series) -> pd.Series:
    text = series.astype(str).str.replace("\u00a0", "", regex=False).str.replace(" ", "", regex=False)
    text = text.str.replace(",", ".", regex=False)
    return pd.to_numeric(text, errors="coerce").fillna(0.0)


def norm_time(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, datetime):
        return f"{value.hour:02d}:{value.minute:02d}"
    s = str(value).strip()
    if not s or s.casefold() == "nan":
        return None
    if " " in s and ":" in s.split(" ", 1)[1]:
        s = s.split(" ", 1)[1]
    m = re.match(r"^(\d{1,2}):(\d{2})", s)
    if m:
        return f"{int(m.group(1)) % 24:02d}:{m.group(2)}"
    try:
        f = float(s.replace(",", "."))
        if 0 <= f < 2:
            total = round((f % 1) * 24 * 60)
            return f"{(total // 60) % 24:02d}:{total % 60:02d}"
    except ValueError:
        pass
    return s


def _base_name(name: str) -> str:
    clean = VERSION_RE.sub("", name)
    return re.sub(r"\s+", " ", Path(clean).with_suffix("").name).strip()


def _version(path: Path) -> int:
    match = VERSION_RE.search(path.name)
    return int(match.group(1)) if match else 0


def _select_sheet(names: list[str], wanted: str) -> str | None:
    if wanted in names:
        return wanted
    for name in names:
        if "реестр" in name.casefold():
            return name
    return None


def _merge_duplicate_alias_columns(df: pd.DataFrame) -> pd.DataFrame:
    for col in list(df.columns):
        if "__dup" not in col:
            continue
        base = col.split("__dup", 1)[0]
        if base in df.columns:
            df[base] = df[base].combine_first(df[col])
        else:
            df = df.rename(columns={col: base})
        if col in df.columns:
            df = df.drop(columns=[col])
    return df


def _unit_from_filename(name: str) -> str | None:
    match = re.search(r"\(([^)]+)\)", name)
    return norm_text(match.group(1)) if match else None


def _clean_text_or_na(value: object):
    text = norm_text(value)
    return text if text and text.casefold() not in {"nan", "<na>", "none"} else pd.NA


def _max_date_iso(series: pd.Series) -> str | None:
    parsed = pd.to_datetime(series, errors="coerce")
    if parsed.notna().any():
        return parsed.max().date().isoformat()
    return None


def _freshness_target(reports: Iterable[FileReport], freshness_date: date | None = None) -> date:
    if freshness_date is not None:
        return freshness_date
    dates = [
        datetime.strptime(report.max_fact_date, "%Y-%m-%d").date()
        for report in reports
        if report.max_fact_date
    ]
    return max(dates) if dates else date.today()


def _volume_sum(df: pd.DataFrame) -> float:
    return float(pd.to_numeric(df["Обьем работ, м3"], errors="coerce").fillna(0).sum())


def _volumes_by_unit(df: pd.DataFrame) -> dict[str, float]:
    if df.empty:
        return {}
    grouped = df.groupby("Подразделение", dropna=False)["Обьем работ, м3"].sum()
    return {str(k): float(v) for k, v in grouped.items()}


def _order_columns(df: pd.DataFrame) -> pd.DataFrame:
    extras = [col for col in df.columns if col not in ALL_COLS and col not in EXCLUDED_OUTPUT_COLS]
    return df[[col for col in OUTPUT_COLS if col in df.columns] + extras]


def _format_sheet(ws) -> None:
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for col_idx, column in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max(12, min(42, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _format_date_columns(ws, names: tuple[str, ...]) -> None:
    header = {cell.value: cell.column for cell in ws[1]}
    for name in names:
        col = header.get(name)
        if not col:
            continue
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 13
        for cell in ws[letter][1:]:
            if cell.value is not None:
                cell.number_format = EXCEL_DATE_FORMAT
